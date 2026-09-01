#!/usr/bin/env python3
"""
STG17 · The baseline and end-line assessment, as an Excel workbook.

    python tools/build_pretest.py
    python tools/build_pretest.py --out "C:/AfDB/Momar/STG17/Workshop Sep 2026"

Produces two files:

    STG17_Pre-test_Initial_Knowledge.xlsx   the participant workbook, EN and FR sheets
    STG17_Assessment_Answer_Key.xlsx        the answer key — facilitators only

--------------------------------------------------------------------------
WHY THE INSTRUMENT HAS TWO PARTS
--------------------------------------------------------------------------
Part A asks participants to rate their own confidence. Part B tests what they
actually know. Both are needed, and running only the first is the common mistake.

Self-rated confidence often *falls* between a baseline and an end-line, because a
participant who arrives not knowing what RAG is cannot know how little they know,
and five days later they can. An office running only Part A would report the
workshop as a failure. Part B moves in the direction the training actually
produced, and the gap between the two is itself a finding worth reporting.

The same workbook is used before and after — one dropdown switches the moment —
so the two measurements are strictly comparable. Different instruments before and
after measure the instrument, not the training.

--------------------------------------------------------------------------
WHY THE ANSWER KEY IS A SEPARATE FILE
--------------------------------------------------------------------------
A hidden worksheet is unhidden in two clicks. A pre-test whose key travelled with
it measures nothing at all, so the key never enters the participant workbook.
"""

from __future__ import annotations

import argparse
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent.parent

NAVY, JADE, AMBER = "0B2545", "1B7A43", "F2A900"
INK, MUTED, LINE, WASH, WHITE = "33403A", "6B7B75", "D5E6DF", "F4F8F5", "FFFFFF"

THIN = Side(style="thin", color=LINE)
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


# ---------------------------------------------------------------------------
#  Part A — self-assessed capability. Ten statements, one per thread of the week.
# ---------------------------------------------------------------------------
SKILLS = [
    ("D1", "Explain to a colleague the difference between an LLM, retrieval-augmented "
           "generation (RAG), fine-tuning and an agent.",
           "Expliquer à un collègue la différence entre un LLM, la génération augmentée par "
           "récupération (RAG), l'affinage et un agent."),
    ("D1", "Decide whether a given task needs prompt engineering, RAG, or fine-tuning.",
           "Décider si une tâche donnée relève de l'ingénierie de prompt, du RAG ou de l'affinage."),
    ("D1", "Estimate the memory and hardware an open-weight model needs before buying anything.",
           "Estimer la mémoire et le matériel qu'exige un modèle à poids ouverts avant tout achat."),
    ("D1", "Decide whether a use case may run on a cloud API or must stay inside the office.",
           "Décider si un cas d'usage peut passer par une API en nuage ou doit rester dans l'office."),
    ("D2", "Write a prompt with worked examples and a fixed output format that code can check.",
           "Rédiger un prompt avec des exemples traités et un format de sortie vérifiable par du code."),
    ("D2", "Build an evaluation set and measure whether a change to a prompt improved anything.",
           "Construire un jeu d'évaluation et mesurer si une modification de prompt a amélioré quelque chose."),
    ("D3", "Acquire and process an openly licensed non-traditional dataset, respecting its licence.",
           "Acquérir et traiter un jeu de données non traditionnel sous licence ouverte, en respectant sa licence."),
    ("D4", "Compute a subnational indicator from night-time lights for my own country.",
           "Calculer un indicateur infranational à partir des lumières nocturnes pour mon propre pays."),
    ("D4", "Validate a proxy indicator against official statistics and state its limitations honestly.",
           "Valider un indicateur indirect contre des statistiques officielles et en énoncer honnêtement les limites."),
    ("D5", "Publish a reproducible analysis on GitHub with its licence, metadata and documentation.",
           "Publier une analyse reproductible sur GitHub avec sa licence, ses métadonnées et sa documentation."),
]

SCALE_EN = ["1 — I could not begin", "2 — I know the words only", "3 — With help",
            "4 — On my own", "5 — I could teach it"]
SCALE_FR = ["1 — Je ne saurais pas commencer", "2 — Je connais les mots seulement",
            "3 — Avec de l'aide", "4 — Seul(e)", "5 — Je pourrais l'enseigner"]


# ---------------------------------------------------------------------------
#  Part B — knowledge. Ten questions, each with one defensible answer.
#  (day, question_en, options_en, question_fr, options_fr, correct letter)
# ---------------------------------------------------------------------------
KNOWLEDGE = [
    ("D1",
     "When a language model produces text, what is it doing?",
     ["It looks the answer up in a database it was given",
      "It predicts the next token, repeatedly, from what came before",
      "It searches the internet and summarises what it finds",
      "It applies rules written by its developers"],
     "Lorsqu'un modèle de langage produit du texte, que fait-il ?",
     ["Il consulte la réponse dans une base de données qu'on lui a fournie",
      "Il prédit le token suivant, de façon répétée, à partir de ce qui précède",
      "Il cherche sur internet et résume ce qu'il trouve",
      "Il applique des règles écrites par ses développeurs"], "B"),

    ("D1",
     "Fine-tuning a model is best suited to teaching it:",
     ["The facts and figures contained in your publications",
      "The style and format in which your office writes",
      "Which documents to retrieve for a question",
      "How to refuse a question it cannot answer"],
     "L'affinage d'un modèle sert avant tout à lui enseigner :",
     ["Les faits et les chiffres contenus dans vos publications",
      "Le style et le format dans lesquels votre office rédige",
      "Quels documents récupérer pour une question donnée",
      "Comment refuser une question à laquelle il ne peut répondre"], "B"),

    ("D1",
     "In a RAG assistant, the retrieved passages do not contain the answer. What should it do?",
     ["Answer from its own general knowledge",
      "Give the closest answer it can, with a caveat",
      "State that the answer is not in the documents",
      "Retrieve more passages until it finds something"],
     "Dans un assistant RAG, les passages récupérés ne contiennent pas la réponse. Que doit-il faire ?",
     ["Répondre à partir de ses connaissances générales",
      "Donner la réponse la plus proche possible, avec une réserve",
      "Déclarer que la réponse ne figure pas dans les documents",
      "Récupérer davantage de passages jusqu'à trouver quelque chose"], "C"),

    ("D1",
     "In an agent, who executes a tool the model has asked for?",
     ["The model, directly",
      "Your own code, which may refuse the request",
      "The model provider, on its servers",
      "The tool itself, automatically"],
     "Dans un agent, qui exécute un outil demandé par le modèle ?",
     ["Le modèle, directement",
      "Votre propre code, qui peut refuser la demande",
      "Le fournisseur du modèle, sur ses serveurs",
      "L'outil lui-même, automatiquement"], "B"),

    ("D1",
     "Roughly how much memory does an 8-billion-parameter model need at 4-bit precision?",
     ["About 0.5 GB", "About 5 GB", "About 32 GB", "About 140 GB"],
     "Quelle mémoire exige approximativement un modèle de 8 milliards de paramètres en 4 bits ?",
     ["Environ 0,5 Go", "Environ 5 Go", "Environ 32 Go", "Environ 140 Go"], "B"),

    # Rewritten: the four options now complete the stem identically. In the first
    # version three began "If ..." and the correct one began "Not without ...",
    # so the option that broke the pattern announced itself as the answer — the
    # question measured pattern recognition rather than knowledge. The jargon
    # ("third-party inference API") also went: a baseline test must be readable
    # by someone who has not yet attended the week that teaches the term.
    ("D1",
     "Your office wants to use an AI service hosted outside the country on data that "
     "identifies individual respondents. When is that permitted?",
     ["When the provider guarantees that it will not keep the data",
      "When a confidentiality agreement has been signed with the provider",
      "When it is only a pilot, on a small sample of records",
      "When the law governing your statistics permits the transfer"],
     "Votre office souhaite utiliser un service d'IA hébergé hors du pays sur des données qui "
     "identifient des répondants. Quand cela est-il permis ?",
     ["Lorsque le fournisseur garantit qu'il ne conservera pas les données",
      "Lorsqu'un accord de confidentialité a été signé avec le fournisseur",
      "Lorsqu'il ne s'agit que d'un pilote, sur un petit échantillon d'enregistrements",
      "Lorsque la loi qui régit vos statistiques autorise ce transfert"], "D"),

    ("D2",
     "Which of these has published evidence of improving results on statistical tasks?",
     ["Adding “You are an expert statistician” to the prompt",
      "Including worked examples in the prompt",
      "Being polite to the model",
      "Asking the model to double-check itself"],
     "Laquelle de ces pratiques a des preuves publiées d'amélioration sur des tâches statistiques ?",
     ["Ajouter « Vous êtes un statisticien expert » au prompt",
      "Inclure des exemples traités dans le prompt",
      "Être poli avec le modèle",
      "Demander au modèle de se relire"], "B"),

    ("D2",
     "Setting the temperature to 0 guarantees:",
     ["The same output every time",
      "A factually correct answer",
      "Reduced variation, but not identical output",
      "That the model will not hallucinate"],
     "Régler la température à 0 garantit :",
     ["La même sortie à chaque fois",
      "Une réponse factuellement exacte",
      "Une variation réduite, mais pas une sortie identique",
      "Que le modèle n'hallucinera pas"], "C"),

    ("D3",
     "A dataset is licensed CC BY-NC-SA 4.0. A national indicator derived from it must:",
     ["Be published under any licence your office chooses",
      "Carry the same NonCommercial and ShareAlike terms",
      "Be kept internal and never published",
      "Only credit the source, with no other obligation"],
     "Un jeu de données est sous licence CC BY-NC-SA 4.0. Un indicateur national qui en dérive doit :",
     ["Être publié sous la licence choisie par votre office",
      "Porter les mêmes clauses NonCommercial et ShareAlike",
      "Rester interne et ne jamais être publié",
      "Seulement créditer la source, sans autre obligation"], "B"),

    ("D4",
     "Night-time lights are most defensible as a proxy for:",
     ["The level of GDP, in absolute terms",
      "Household connection rates to the grid",
      "Where electrification and settlement have spread",
      "Informal economic activity"],
     "Les lumières nocturnes sont le plus défendables comme indicateur indirect :",
     ["Du niveau du PIB, en valeur absolue",
      "Du taux de raccordement des ménages au réseau",
      "De l'extension de l'électrification et de l'habitat",
      "De l'activité économique informelle"], "C"),
]

TEXT = {
    "en": {
        "sheet": "Pre-test EN",
        "title": "Pre-test — Initial Knowledge",
        "sub": "Baseline and end-line self-assessment · STG17 Action Plan activity 4.1.1",
        "workshop": "STG17 Technical Workshop · Emerging Issues, Emerging Practice · "
                    "Kigali, 28 September – 2 October 2026",
        "howto": "This questionnaire is completed TWICE — once before the workshop opens and "
                 "again after it closes — with the same questions both times, so that the two "
                 "measurements can be compared. It takes about 12 minutes. It is not an "
                 "examination: nobody is graded, and an honest baseline is what makes the "
                 "end-line meaningful. Please answer without looking anything up.",
        "who": "About you",
        "fields": ["Full name", "Country", "Institution", "Unit or department",
                   "Email", "Date (dd/mm/yyyy)", "Moment"],
        "moment": ["Pre-test (before the workshop)", "Post-test (after the workshop)"],
        "partA": "Part A · What you can do today",
        "partA_lead": "For each statement, choose the number that best describes where you stand "
                      "right now. There is no right answer.",
        "partB": "Part B · What you know today",
        "partB_lead": "One answer per question. If you do not know, choose the answer you think "
                      "most likely — but please do not look it up.",
        "cols_a": ["#", "I could…", "Your rating"],
        "cols_b": ["#", "Question", "Options", "Your answer"],
        "scale": "Scale:  " + "   ·   ".join(SCALE_EN),
        "thanks": "Thank you. Please save this file as  PRETEST_<COUNTRY>_<YOUR NAME>.xlsx  and "
                  "return it to the workshop secretariat.",
        "answer_hint": "Choose A, B, C or D",
    },
    "fr": {
        "sheet": "Pré-test FR",
        "title": "Pré-test — Connaissances initiales",
        "sub": "Auto-évaluation initiale et finale · Plan d'action STG17, activité 4.1.1",
        "workshop": "Atelier technique GTS 17 · Enjeux émergents, pratiques émergentes · "
                    "Kigali, 28 septembre – 2 octobre 2026",
        "howto": "Ce questionnaire est rempli DEUX FOIS — une fois avant l'ouverture de "
                 "l'atelier et une fois après sa clôture — avec les mêmes questions, afin que "
                 "les deux mesures soient comparables. Il demande environ 12 minutes. Ce n'est "
                 "pas un examen : personne n'est noté, et c'est un état initial honnête qui "
                 "donne son sens à la mesure finale. Merci de répondre sans rien consulter.",
        "who": "À propos de vous",
        "fields": ["Nom et prénom", "Pays", "Institution", "Unité ou département",
                   "Courriel", "Date (jj/mm/aaaa)", "Moment"],
        "moment": ["Pré-test (avant l'atelier)", "Post-test (après l'atelier)"],
        "partA": "Partie A · Ce que vous savez faire aujourd'hui",
        "partA_lead": "Pour chaque énoncé, choisissez le chiffre qui décrit le mieux votre "
                      "situation actuelle. Il n'y a pas de bonne réponse.",
        "partB": "Partie B · Ce que vous savez aujourd'hui",
        "partB_lead": "Une seule réponse par question. Si vous ne savez pas, choisissez la "
                      "réponse qui vous semble la plus probable — mais ne la cherchez pas.",
        "cols_a": ["#", "Je saurais…", "Votre note"],
        "cols_b": ["#", "Question", "Options", "Votre réponse"],
        "scale": "Échelle :  " + "   ·   ".join(SCALE_FR),
        "thanks": "Merci. Enregistrez ce fichier sous  PRETEST_<PAYS>_<VOTRE NOM>.xlsx  et "
                  "retournez-le au secrétariat de l'atelier.",
        "answer_hint": "Choisissez A, B, C ou D",
    },
}


# ---------------------------------------------------------------------------
#  Styling helpers
# ---------------------------------------------------------------------------
def band(ws, row, col_from, col_to, text, *, fill, colour, size, bold=True, height=None,
         align="left", wrap=False, italic=False):
    ws.merge_cells(start_row=row, start_column=col_from, end_row=row, end_column=col_to)
    cell = ws.cell(row=row, column=col_from, value=text)
    cell.fill = PatternFill("solid", fgColor=fill)
    cell.font = Font(name="Calibri", size=size, bold=bold, italic=italic, color=colour)
    cell.alignment = Alignment(horizontal=align, vertical="center", wrap_text=wrap,
                               indent=1 if align == "left" else 0)
    if height:
        ws.row_dimensions[row].height = height
    return cell


def build_sheet(ws, lang: str) -> None:
    t = TEXT[lang]
    skills_col = 1 if lang == "en" else 2   # index into (day, en, fr) minus the day
    widths = [5, 52, 62, 17]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.sheet_view.showGridLines = False

    # --- masthead ---------------------------------------------------------
    band(ws, 1, 1, 4, t["title"], fill=NAVY, colour=WHITE, size=20, height=38)
    band(ws, 2, 1, 4, t["sub"], fill=NAVY, colour="C9DAD2", size=10, bold=False, height=20)
    band(ws, 3, 1, 4, t["workshop"], fill=AMBER, colour=NAVY, size=9.5, height=18)
    band(ws, 4, 1, 4, t["howto"], fill=WASH, colour=INK, size=10, bold=False,
         height=54, wrap=True)

    # --- who ---------------------------------------------------------------
    row = 6
    band(ws, row, 1, 4, t["who"], fill=JADE, colour=WHITE, size=12, height=22)
    row += 1
    first_field_row = row
    for label in t["fields"]:
        ws.cell(row=row, column=1, value="").border = BOX
        c = ws.cell(row=row, column=2, value=label)
        c.font = Font(name="Calibri", size=10, bold=True, color=INK)
        c.alignment = Alignment(vertical="center", indent=1)
        c.fill = PatternFill("solid", fgColor=WASH)
        c.border = BOX
        entry = ws.cell(row=row, column=3, value="")
        entry.border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
        entry.fill = PatternFill("solid", fgColor=WHITE)
        ws.cell(row=row, column=4, value="").border = BOX
        ws.row_dimensions[row].height = 19
        row += 1

    moment_dv = DataValidation(type="list", formula1='"{}"'.format(",".join(t["moment"])),
                               allow_blank=True, showDropDown=False)
    ws.add_data_validation(moment_dv)
    moment_dv.add(ws.cell(row=first_field_row + len(t["fields"]) - 1, column=3))

    # --- Part A ------------------------------------------------------------
    row += 1
    band(ws, row, 1, 4, t["partA"], fill=NAVY, colour=WHITE, size=13, height=24)
    row += 1
    band(ws, row, 1, 4, t["partA_lead"], fill=WASH, colour=MUTED, size=9.5, bold=False,
         italic=True, height=17)
    row += 1
    band(ws, row, 1, 4, t["scale"], fill=WHITE, colour=JADE, size=9, height=17)
    row += 1

    for j, header in enumerate(t["cols_a"], start=1):
        span = (1, 1) if j == 1 else ((2, 3) if j == 2 else (4, 4))
        band(ws, row, span[0], span[1], header, fill=JADE, colour=WHITE, size=10,
             height=20, align="center" if j != 2 else "left")
    row += 1

    rating_dv = DataValidation(type="list", formula1='"1,2,3,4,5"', allow_blank=True,
                               showDropDown=False,
                               errorTitle="1 to 5", error="Please choose a whole number from 1 to 5.")
    ws.add_data_validation(rating_dv)

    part_a_first = row
    for n, item in enumerate(SKILLS, start=1):
        text = item[skills_col]
        fill = WHITE if n % 2 else WASH
        num = ws.cell(row=row, column=1, value=f"A{n}")
        num.font = Font(name="Calibri", size=10, bold=True, color=JADE)
        num.alignment = Alignment(horizontal="center", vertical="center")
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=3)
        body = ws.cell(row=row, column=2, value=text)
        body.font = Font(name="Calibri", size=10, color=INK)
        body.alignment = Alignment(vertical="center", wrap_text=True, indent=1)
        ans = ws.cell(row=row, column=4, value="")
        ans.alignment = Alignment(horizontal="center", vertical="center")
        ans.font = Font(name="Calibri", size=12, bold=True, color=NAVY)
        rating_dv.add(ans)
        for col in range(1, 5):
            cell = ws.cell(row=row, column=col)
            cell.border = BOX
            if cell.fill.fgColor.rgb in (None, "00000000"):
                cell.fill = PatternFill("solid", fgColor=fill)
        ans.fill = PatternFill("solid", fgColor=WHITE)
        ws.row_dimensions[row].height = 30
        row += 1
    part_a_last = row - 1

    # --- Part B ------------------------------------------------------------
    row += 1
    band(ws, row, 1, 4, t["partB"], fill=NAVY, colour=WHITE, size=13, height=24)
    row += 1
    band(ws, row, 1, 4, t["partB_lead"], fill=WASH, colour=MUTED, size=9.5, bold=False,
         italic=True, height=17)
    row += 1
    for j, header in enumerate(t["cols_b"], start=1):
        band(ws, row, j, j, header, fill=JADE, colour=WHITE, size=10, height=20,
             align="center" if j in (1, 4) else "left")
    row += 1

    letter_dv = DataValidation(type="list", formula1='"A,B,C,D"', allow_blank=True,
                               showDropDown=False,
                               errorTitle="A to D", error="Please choose A, B, C or D.")
    ws.add_data_validation(letter_dv)

    part_b_first = row
    for n, (day, q_en, o_en, q_fr, o_fr, _correct) in enumerate(KNOWLEDGE, start=1):
        question = q_en if lang == "en" else q_fr
        options = o_en if lang == "en" else o_fr
        fill = WHITE if n % 2 else WASH

        num = ws.cell(row=row, column=1, value=f"B{n}")
        num.font = Font(name="Calibri", size=10, bold=True, color=JADE)
        num.alignment = Alignment(horizontal="center", vertical="center")

        q = ws.cell(row=row, column=2, value=question)
        q.font = Font(name="Calibri", size=10, bold=True, color=NAVY)
        q.alignment = Alignment(vertical="top", wrap_text=True, indent=1)

        opts = ws.cell(row=row, column=3,
                       value="\n".join(f"{chr(65 + i)}.  {o}" for i, o in enumerate(options)))
        opts.font = Font(name="Calibri", size=9.5, color=INK)
        opts.alignment = Alignment(vertical="top", wrap_text=True, indent=1)

        ans = ws.cell(row=row, column=4, value="")
        ans.alignment = Alignment(horizontal="center", vertical="center")
        ans.font = Font(name="Calibri", size=12, bold=True, color=NAVY)
        letter_dv.add(ans)

        for col in range(1, 5):
            cell = ws.cell(row=row, column=col)
            cell.border = BOX
            cell.fill = PatternFill("solid", fgColor=fill)
        ans.fill = PatternFill("solid", fgColor=WHITE)
        ws.row_dimensions[row].height = 62
        row += 1
    part_b_last = row - 1

    # --- close -------------------------------------------------------------
    row += 1
    band(ws, row, 1, 4, t["thanks"], fill=AMBER, colour=NAVY, size=10, height=26, wrap=True)

    ws.freeze_panes = "A6"
    ws.print_title_rows = "1:3"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    return part_a_first, part_a_last, part_b_first, part_b_last


def build_key() -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = "Answer key"
    ws.sheet_view.showGridLines = False
    for col, w in zip("ABCDE", (7, 58, 12, 12, 40)):
        ws.column_dimensions[col].width = w

    band(ws, 1, 1, 5, "STG17 — Assessment answer key", fill=NAVY, colour=WHITE, size=18, height=34)
    band(ws, 2, 1, 5, "FACILITATORS ONLY — do not distribute to participants before the end-line",
         fill="C0392B", colour=WHITE, size=11, height=22)
    band(ws, 3, 1, 5,
         "Part A is self-assessment and has no correct answer: score it as the mean of the ten "
         "ratings, from 1 to 5. Part B is scored out of 10. Report both, and report them "
         "separately — a confidence score that falls while the knowledge score rises is the "
         "normal, healthy pattern, not a failure.",
         fill=WASH, colour=INK, size=10, bold=False, height=48, wrap=True)

    row = 5
    for j, h in enumerate(["#", "Question (EN)", "Correct", "Session", "What it tests"], start=1):
        band(ws, row, j, j, h, fill=JADE, colour=WHITE, size=10, height=20,
             align="center" if j in (1, 3, 4) else "left")
    row += 1

    tests = [
        "The generation mechanism — no lookup, no database",
        "Fine-tuning teaches style, not facts",
        "Correct refusal when the corpus cannot answer",
        "Where the approval gate sits in an agent",
        "Memory arithmetic: parameters × bytes per parameter",
        "Microdata leaving the country needs a legal basis, not a promise",
        "Evidence for worked examples; none for personas",
        "Temperature 0 is not reproducibility",
        "CC BY-NC-SA terms propagate into derivatives",
        "What NTL proxies well, and what it does not",
    ]
    for n, ((day, q_en, o_en, _q, _o, correct), what) in enumerate(zip(KNOWLEDGE, tests), start=1):
        for col, value in enumerate([f"B{n}", q_en, correct, day, what], start=1):
            cell = ws.cell(row=row, column=col, value=value)
            cell.border = BOX
            cell.fill = PatternFill("solid", fgColor=WHITE if n % 2 else WASH)
            cell.alignment = Alignment(vertical="center", wrap_text=col in (2, 5), indent=1)
            cell.font = Font(name="Calibri", size=10,
                             bold=col in (1, 3),
                             color=JADE if col == 3 else INK)
        ws.row_dimensions[row].height = 30
        row += 1

    row += 1
    band(ws, row, 1, 5,
         "Impact = post-test minus pre-test, computed per participant and reported as a "
         "distribution rather than a mean. Feeds STG17 Action Plan activity 4.1.1 "
         "(competencies reference framework).",
         fill=AMBER, colour=NAVY, size=10, height=30, wrap=True)
    return wb


def main() -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")

    ap = argparse.ArgumentParser(description="Build the STG17 baseline/end-line assessment.")
    ap.add_argument("--out", default=str(ROOT / "concept-note"))
    args = ap.parse_args()
    out = Path(args.out)
    out.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    wb.remove(wb.active)
    for lang in ("en", "fr"):
        ws = wb.create_sheet(TEXT[lang]["sheet"])
        build_sheet(ws, lang)
    wb.properties.title = "STG17 — Pre-test: Initial Knowledge"
    wb.properties.creator = "African Development Bank · AU STATAFRIC · STG17"

    failed = []
    for workbook, name in ((wb, "STG17_Pre-test_Initial_Knowledge.xlsx"),
                           (build_key(), "STG17_Assessment_Answer_Key.xlsx")):
        path = out / name
        try:
            workbook.save(path)
        except PermissionError:
            failed.append(path)
            print(f"[!] {path.name} is open in another program — not written.")
            continue
        print(f"Written: {path}")

    print(f"  {len(SKILLS)} self-assessment statements + {len(KNOWLEDGE)} knowledge questions, "
          f"in English and French.")
    return 1 if failed else 0


if __name__ == "__main__":
    raise SystemExit(main())
